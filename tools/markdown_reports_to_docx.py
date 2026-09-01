from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PAGE_MARGIN_INCHES = 0.75
BODY_FONT = "Arial"
BODY_SIZE = Pt(10)
HEADING_COLOR = RGBColor(31, 78, 121)
BODY_COLOR = RGBColor(0, 0, 0)
TABLE_HEADER_FILL = "D9EAF7"
TABLE_BORDER_COLOR = "A6A6A6"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_borders(cell, color: str = TABLE_BORDER_COLOR, size: str = "6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, margin_dxa: int = 90) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side in ("top", "start", "bottom", "end"):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(margin_dxa))
        node.set(qn("w:type"), "dxa")


def apply_run_formatting(paragraph, text: str) -> None:
    pattern = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*)")
    cursor = 0
    for match in pattern.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor : match.start()])
            run.font.name = BODY_FONT
            run.font.size = BODY_SIZE
            run.font.color.rgb = BODY_COLOR
        token = match.group(0)
        if token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Courier New"
            run.font.size = Pt(9)
            run.font.color.rgb = BODY_COLOR
        elif token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
            run.font.name = BODY_FONT
            run.font.size = BODY_SIZE
            run.font.color.rgb = BODY_COLOR
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        run.font.name = BODY_FONT
        run.font.size = BODY_SIZE
        run.font.color.rgb = BODY_COLOR


def is_markdown_table_start(lines: list[str], idx: int) -> bool:
    if idx + 1 >= len(lines):
        return False
    if not lines[idx].strip().startswith("|"):
        return False
    separator_cells = [
        cell.strip()
        for cell in lines[idx + 1].strip().strip("|").split("|")
    ]
    return bool(separator_cells) and all(
        re.fullmatch(r":?-{3,}:?", cell) for cell in separator_cells
    )


def parse_table(lines: list[str], idx: int) -> tuple[list[list[str]], int]:
    table_lines = []
    while idx < len(lines) and lines[idx].strip().startswith("|"):
        table_lines.append(lines[idx].strip())
        idx += 1

    rows: list[list[str]] = []
    for line_no, line in enumerate(table_lines):
        if line_no == 1:
            continue
        cells = [cell.strip() for cell in line.strip("|").split("|")]
        rows.append(cells)
    return rows, idx


def add_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    column_count = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=column_count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.allow_autofit = True

    for row_idx, row in enumerate(rows):
        for col_idx in range(column_count):
            cell = table.cell(row_idx, col_idx)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_borders(cell)
            set_cell_margins(cell)
            text = row[col_idx] if col_idx < len(row) else ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.05
            apply_run_formatting(paragraph, text)
            for run in paragraph.runs:
                run.font.name = BODY_FONT
                run.font.size = Pt(8.5)
                run.font.color.rgb = BODY_COLOR
                if row_idx == 0:
                    run.bold = True
            if row_idx == 0:
                set_cell_shading(cell, TABLE_HEADER_FILL)

    document.add_paragraph()


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(PAGE_MARGIN_INCHES)
    section.bottom_margin = Inches(PAGE_MARGIN_INCHES)
    section.left_margin = Inches(PAGE_MARGIN_INCHES)
    section.right_margin = Inches(PAGE_MARGIN_INCHES)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = BODY_SIZE
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    for name, size in [
        ("Heading 1", 16),
        ("Heading 2", 13),
        ("Heading 3", 11),
    ]:
        style = styles[name]
        style.font.name = BODY_FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = HEADING_COLOR
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(4)


def add_markdown_line(document: Document, line: str) -> None:
    stripped = line.strip()
    if not stripped:
        return

    heading_match = re.match(r"^(#{1,6})\s+(.*)$", stripped)
    if heading_match:
        level = min(len(heading_match.group(1)), 3)
        text = heading_match.group(2)
        paragraph = document.add_heading(level=level)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        apply_run_formatting(paragraph, text)
        for run in paragraph.runs:
            run.font.color.rgb = HEADING_COLOR
            run.bold = True
        return

    if stripped.startswith("- "):
        paragraph = document.add_paragraph(style="List Bullet")
        apply_run_formatting(paragraph, stripped[2:])
        return

    numbered = re.match(r"^\d+\.\s+(.*)$", stripped)
    if numbered:
        paragraph = document.add_paragraph()
        apply_run_formatting(paragraph, stripped)
        return

    paragraph = document.add_paragraph()
    apply_run_formatting(paragraph, stripped)


def convert_markdown_to_docx(markdown_path: Path, docx_path: Path) -> None:
    document = Document()
    configure_document(document)

    lines = markdown_path.read_text(encoding="utf-8").splitlines()
    idx = 0
    in_code_block = False
    while idx < len(lines):
        if lines[idx].strip().startswith("```"):
            in_code_block = not in_code_block
            idx += 1
            continue
        if in_code_block:
            paragraph = document.add_paragraph()
            run = paragraph.add_run(lines[idx])
            run.font.name = "Courier New"
            run.font.size = Pt(9)
            run.font.color.rgb = BODY_COLOR
            idx += 1
            continue
        if is_markdown_table_start(lines, idx):
            rows, idx = parse_table(lines, idx)
            add_table(document, rows)
            continue
        add_markdown_line(document, lines[idx])
        idx += 1

    document.save(docx_path)


def main() -> None:
    if len(sys.argv) < 3 or len(sys.argv) % 2 == 0:
        raise SystemExit(
            "Usage: markdown_reports_to_docx.py input1.md output1.docx [input2.md output2.docx ...]"
        )

    pairs = zip(sys.argv[1::2], sys.argv[2::2])
    for input_path, output_path in pairs:
        convert_markdown_to_docx(Path(input_path), Path(output_path))
        print(f"Saved {output_path}")


if __name__ == "__main__":
    main()
